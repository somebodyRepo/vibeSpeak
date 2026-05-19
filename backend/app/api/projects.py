import json
import re
import zipfile
import io
from datetime import datetime
from urllib.parse import quote
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import desc, func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.database import OutlineDB, ProjectDB, InterviewSessionDB, get_db
from app.models.schemas import (
    OutlineContent,
    OutlineResponse,
    OutlineSection,
    ProjectCreate,
    ProjectListResponse,
    ProjectResponse,
    ProjectUpdate,
    ProjectWithOutlineImport,
    SessionResponse,
)

router = APIRouter(prefix="/projects", tags=["projects"])


@router.post("/", response_model=ProjectResponse)
async def create_project(
    request: ProjectCreate,
    db: AsyncSession = Depends(get_db),
):
    """创建项目"""
    project_id = str(uuid4())

    # 如果指定了 outline_id，验证提纲是否存在
    if request.outline_id:
        result = await db.execute(
            select(OutlineDB).where(OutlineDB.id == request.outline_id)
        )
        if not result.scalar_one_or_none():
            raise HTTPException(status_code=404, detail="Outline not found")

    project = ProjectDB(
        id=project_id,
        name=request.name,
        description=request.description or "",
        outline_id=request.outline_id,
    )
    db.add(project)
    await db.commit()
    await db.refresh(project)

    return ProjectResponse(
        id=project.id,
        name=project.name,
        description=project.description,
        outline_id=project.outline_id,
        table_structure_prompt=project.table_structure_prompt or "",
        summary_table=project.summary_table or "",
        session_count=0,
        created_at=project.created_at,
        updated_at=project.updated_at,
    )


@router.get("/", response_model=ProjectListResponse)
async def list_projects(
    limit: int = 50,
    offset: int = 0,
    db: AsyncSession = Depends(get_db),
):
    """获取项目列表"""
    # 子查询获取会话数量
    session_count_subq = (
        select(
            InterviewSessionDB.project_id,
            func.count(InterviewSessionDB.id).label("count")
        )
        .group_by(InterviewSessionDB.project_id)
        .subquery()
    )

    # 主查询
    query = (
        select(ProjectDB, func.coalesce(session_count_subq.c.count, 0).label("session_count"))
        .outerjoin(session_count_subq, ProjectDB.id == session_count_subq.c.project_id)
        .order_by(desc(ProjectDB.created_at))
        .offset(offset)
        .limit(limit)
    )

    result = await db.execute(query)
    rows = result.all()

    # 获取总数
    count_result = await db.execute(select(ProjectDB))
    total = len(count_result.scalars().all())

    projects = []
    for project, session_count in rows:
        # 获取提纲信息
        outline = None
        if project.outline_id:
            outline_result = await db.execute(
                select(OutlineDB).where(OutlineDB.id == project.outline_id)
            )
            outline_db = outline_result.scalar_one_or_none()
            if outline_db:
                outline = OutlineResponse(
                    id=outline_db.id,
                    name=outline_db.name,
                    content=OutlineContent.model_validate_json(outline_db.content),
                    created_at=outline_db.created_at,
                    updated_at=outline_db.updated_at,
                )

        projects.append(ProjectResponse(
            id=project.id,
            name=project.name,
            description=project.description,
            outline_id=project.outline_id,
            table_structure_prompt=project.table_structure_prompt or "",
            summary_table=project.summary_table or "",
            outline=outline,
            session_count=session_count,
            created_at=project.created_at,
            updated_at=project.updated_at,
        ))

    return ProjectListResponse(projects=projects, total=total)


@router.get("/{project_id}", response_model=ProjectResponse)
async def get_project(
    project_id: str,
    db: AsyncSession = Depends(get_db),
):
    """获取项目详情"""
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # 获取会话数量
    session_result = await db.execute(
        select(func.count(InterviewSessionDB.id))
        .where(InterviewSessionDB.project_id == project_id)
    )
    session_count = session_result.scalar() or 0

    # 获取提纲信息
    outline = None
    if project.outline_id:
        outline_result = await db.execute(
            select(OutlineDB).where(OutlineDB.id == project.outline_id)
        )
        outline_db = outline_result.scalar_one_or_none()
        if outline_db:
            outline = OutlineResponse(
                id=outline_db.id,
                name=outline_db.name,
                content=OutlineContent.model_validate_json(outline_db.content),
                created_at=outline_db.created_at,
                updated_at=outline_db.updated_at,
            )

    return ProjectResponse(
        id=project.id,
        name=project.name,
        description=project.description,
        outline_id=project.outline_id,
        table_structure_prompt=project.table_structure_prompt or "",
        summary_table=project.summary_table or "",
        outline=outline,
        session_count=session_count,
        created_at=project.created_at,
        updated_at=project.updated_at,
    )


@router.put("/{project_id}", response_model=ProjectResponse)
async def update_project(
    project_id: str,
    request: ProjectUpdate,
    db: AsyncSession = Depends(get_db),
):
    """更新项目"""
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    if request.name is not None:
        project.name = request.name

    if request.description is not None:
        project.description = request.description

    if request.outline_id is not None:
        # 验证提纲是否存在
        if request.outline_id:
            outline_result = await db.execute(
                select(OutlineDB).where(OutlineDB.id == request.outline_id)
            )
            if not outline_result.scalar_one_or_none():
                raise HTTPException(status_code=404, detail="Outline not found")
        project.outline_id = request.outline_id

    await db.commit()
    await db.refresh(project)

    # 获取会话数量
    session_result = await db.execute(
        select(func.count(InterviewSessionDB.id))
        .where(InterviewSessionDB.project_id == project_id)
    )
    session_count = session_result.scalar() or 0

    # 获取提纲信息
    outline = None
    if project.outline_id:
        outline_result = await db.execute(
            select(OutlineDB).where(OutlineDB.id == project.outline_id)
        )
        outline_db = outline_result.scalar_one_or_none()
        if outline_db:
            outline = OutlineResponse(
                id=outline_db.id,
                name=outline_db.name,
                content=OutlineContent.model_validate_json(outline_db.content),
                created_at=outline_db.created_at,
                updated_at=outline_db.updated_at,
            )

    return ProjectResponse(
        id=project.id,
        name=project.name,
        description=project.description,
        outline_id=project.outline_id,
        table_structure_prompt=project.table_structure_prompt or "",
        summary_table=project.summary_table or "",
        outline=outline,
        session_count=session_count,
        created_at=project.created_at,
        updated_at=project.updated_at,
    )


@router.delete("/{project_id}")
async def delete_project(
    project_id: str,
    db: AsyncSession = Depends(get_db),
):
    """删除项目 (同时删除关联的会话)"""
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # 删除关联的会话
    session_result = await db.execute(
        select(InterviewSessionDB).where(InterviewSessionDB.project_id == project_id)
    )
    sessions = session_result.scalars().all()
    for session in sessions:
        await db.delete(session)

    await db.delete(project)
    await db.commit()

    return {"success": True, "message": "Project and related sessions deleted"}


# ===== 表格结构提示词 API =====


@router.post("/{project_id}/design-table-structure")
async def design_table_structure(
    project_id: str,
    db: AsyncSession = Depends(get_db),
):
    """触发 LLM 生成表格结构提示词"""
    from app.services.llm_service import llm_service

    # 获取项目
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # 检查是否有提纲
    if not project.outline_id:
        raise HTTPException(status_code=400, detail="项目无关联提纲，无法设计表格结构")

    # 获取提纲
    outline_result = await db.execute(
        select(OutlineDB).where(OutlineDB.id == project.outline_id)
    )
    outline_db = outline_result.scalar_one_or_none()

    if not outline_db:
        raise HTTPException(status_code=400, detail="提纲不存在")

    # 获取已完成的访谈（status='done'）
    session_result = await db.execute(
        select(InterviewSessionDB)
        .where(InterviewSessionDB.project_id == project_id)
        .where(InterviewSessionDB.status == "done")
        .where(InterviewSessionDB.final_content != "")
    )
    sessions = session_result.scalars().all()

    if not sessions:
        raise HTTPException(status_code=400, detail="无已完成的访谈内容，无法设计表格结构")

    # 解析提纲
    outline_content = OutlineContent.model_validate_json(outline_db.content)

    # 收集访谈内容
    final_contents = [s.final_content for s in sessions if s.final_content]

    # 调用 LLM 设计提示词
    prompt = await llm_service.design_table_structure_prompt(
        outline=outline_content.model_dump(),
        final_contents=final_contents,
    )

    return {
        "success": True,
        "prompt": prompt,
        "message": "表格结构提示词生成成功",
    }


class TableStructureUpdate(BaseModel):
    prompt: str


@router.put("/{project_id}/table-structure")
async def update_table_structure(
    project_id: str,
    request: TableStructureUpdate,
    db: AsyncSession = Depends(get_db),
):
    """保存表格结构提示词"""
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    project.table_structure_prompt = request.prompt
    await db.commit()

    return {
        "success": True,
        "message": "表格结构提示词保存成功",
    }


@router.get("/{project_id}/table-structure")
async def get_table_structure(
    project_id: str,
    db: AsyncSession = Depends(get_db),
):
    """获取当前保存的表格结构提示词"""
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    return {
        "success": True,
        "prompt": project.table_structure_prompt or "",
        "has_outline": bool(project.outline_id),
    }


# ===== 批量表格生成 API =====


@router.post("/{project_id}/generate-all-tables")
async def generate_all_tables(
    project_id: str,
    db: AsyncSession = Depends(get_db),
):
    """触发批量生成所有访谈表格"""
    from app.services.batch_table_service import batch_table_service

    # 获取项目
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    if not project.table_structure_prompt:
        raise HTTPException(status_code=400, detail="项目未设置表格结构提示词，请先设计表格结构")

    # 获取需要生成表格的会话（done状态且未生成表格）
    session_result = await db.execute(
        select(InterviewSessionDB)
        .where(InterviewSessionDB.project_id == project_id)
        .where(InterviewSessionDB.status == "done")
        .where(InterviewSessionDB.final_content != "")
        .where(InterviewSessionDB.table_content == "")
    )
    sessions = session_result.scalars().all()

    if not sessions:
        return {
            "success": True,
            "message": "没有需要生成表格的会话",
            "total_count": 0,
            "pending_count": 0,
        }

    # 准备会话内容
    session_contents = {s.id: s.final_content for s in sessions}
    session_ids = [s.id for s in sessions]

    # 启动批量生成
    progress = batch_table_service.start_batch_generation(
        project_id=project_id,
        session_ids=session_ids,
        table_structure_prompt=project.table_structure_prompt,
        session_contents=session_contents,
    )

    return {
        "success": True,
        "message": f"开始批量生成 {len(sessions)} 个表格",
        "total_count": progress.total_count,
        "pending_count": progress.total_count - progress.completed_count - progress.error_count,
    }


@router.get("/{project_id}/table-generation-progress")
async def get_table_generation_progress(
    project_id: str,
    db: AsyncSession = Depends(get_db),
):
    """获取批量表格生成进度"""
    from app.services.batch_table_service import batch_table_service

    # 验证项目存在
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    progress = batch_table_service.get_progress(project_id)

    if not progress:
        # 返回默认进度
        return {
            "success": True,
            "is_running": False,
            "total_count": 0,
            "completed_count": 0,
            "error_count": 0,
            "tasks": [],
        }

    # 构建任务状态列表
    tasks = []
    for session_id, task in progress.tasks.items():
        tasks.append({
            "session_id": session_id,
            "status": task.status,
            "error_message": task.error_message,
            "retry_count": task.retry_count,
        })

    return {
        "success": True,
        "is_running": progress.is_running,
        "total_count": progress.total_count,
        "completed_count": progress.completed_count,
        "error_count": progress.error_count,
        "tasks": tasks,
    }


# ===== 表格整合汇总 API =====


@router.post("/{project_id}/summarize-tables")
async def summarize_tables(
    project_id: str,
    db: AsyncSession = Depends(get_db),
):
    """触发整合所有已生成表格的会话，生成汇总表格"""
    from app.services.llm_service import llm_service

    # 获取项目
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    if not project.table_structure_prompt:
        raise HTTPException(status_code=400, detail="项目未设置表格结构提示词")

    # 获取所有 tabled 状态的会话
    session_result = await db.execute(
        select(InterviewSessionDB)
        .where(InterviewSessionDB.project_id == project_id)
        .where(InterviewSessionDB.status == "tabled")
        .where(InterviewSessionDB.table_content != "")
    )
    sessions = session_result.scalars().all()

    if not sessions:
        raise HTTPException(status_code=400, detail="没有已生成表格的会话，请先生成表格")

    # 解析各会话的表格数据
    session_tables = []
    for session in sessions:
        try:
            table_data = json.loads(session.table_content)
            session_tables.append(table_data)
        except json.JSONDecodeError:
            continue

    if not session_tables:
        raise HTTPException(status_code=400, detail="没有有效的表格数据")

    # 调用 LLM 整合汇总
    summary_markdown = await llm_service.summarize_tables(
        table_structure_prompt=project.table_structure_prompt,
        session_tables=session_tables,
    )

    # 保存汇总表格
    project.summary_table = summary_markdown
    await db.commit()

    return {
        "success": True,
        "message": f"成功整合 {len(sessions)} 个访谈的表格",
        "summary_table": summary_markdown,
        "session_count": len(sessions),
    }


@router.get("/{project_id}/summary-table")
async def get_summary_table(
    project_id: str,
    db: AsyncSession = Depends(get_db),
):
    """获取项目的汇总表格"""
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    return {
        "success": True,
        "summary_table": project.summary_table or "",
        "has_prompt": bool(project.table_structure_prompt),
    }


@router.get("/{project_id}/export-summary")
async def export_summary_table(
    project_id: str,
    format: str = "md",  # md or xlsx
    db: AsyncSession = Depends(get_db),
):
    """导出汇总表格"""
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    if not project.summary_table:
        raise HTTPException(status_code=400, detail="项目尚无汇总表格，请先整合汇总")

    # Generate filename
    date_str = datetime.now().strftime("%Y%m%d")
    safe_name = re.sub(r'[^\w一-鿿]', '_', project.name)

    if format == "xlsx":
        # Generate Excel file
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
        except ImportError:
            raise HTTPException(status_code=500, detail="Excel export not available")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "汇总表格"

        # Parse markdown table
        lines = project.summary_table.split('\n')
        table_lines = [l for l in lines if l.strip().startswith('|')]

        if len(table_lines) < 2:
            raise HTTPException(status_code=400, detail="无效的汇总表格格式")

        # Parse header
        header_cells = [c.strip() for c in table_lines[0].split('|') if c.strip()]
        for i, cell in enumerate(header_cells):
            ws.cell(row=1, column=i+1, value=cell)
            ws.cell(row=1, column=i+1).font = Font(bold=True)
            ws.cell(row=1, column=i+1).fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

        # Parse body (skip separator line)
        for row_idx, line in enumerate(table_lines[2:], start=2):
            cells = [c.strip() for c in line.split('|') if c.strip()]
            for col_idx, cell in enumerate(cells, start=1):
                ws.cell(row=row_idx, column=col_idx, value=cell)

        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        # Save to bytes
        from io import BytesIO
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f"{safe_name}_汇总_{date_str}.xlsx"
        encoded_filename = quote(filename)

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
            }
        )
    else:
        # Export as Markdown
        filename = f"{safe_name}_汇总_{date_str}.md"
        encoded_filename = quote(filename)

        return StreamingResponse(
            iter([project.summary_table.encode()]),
            media_type="text/markdown",
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
            }
        )


@router.get("/{project_id}/export")
async def export_project(
    project_id: str,
    db: AsyncSession = Depends(get_db),
):
    """导出项目所有内容"""
    result = await db.execute(
        select(ProjectDB).where(ProjectDB.id == project_id)
    )
    project = result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # 获取所有会话
    session_result = await db.execute(
        select(InterviewSessionDB)
        .where(InterviewSessionDB.project_id == project_id)
        .order_by(InterviewSessionDB.created_at)
    )
    sessions = session_result.scalars().all()

    # 创建 ZIP 文件
    zip_buffer = io.BytesIO()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        # 添加项目信息
        project_info = {
            "id": project.id,
            "name": project.name,
            "description": project.description,
            "created_at": project.created_at.isoformat(),
            "sessions": [
                {
                    "id": s.id,
                    "filename": s.filename,
                    "status": s.status,
                    "created_at": s.created_at.isoformat(),
                }
                for s in sessions
            ],
        }
        zip_file.writestr(
            f"{project.name}/project_info.json",
            json.dumps(project_info, ensure_ascii=False, indent=2)
        )

        # 添加每个会话的内容
        for session in sessions:
            session_dir = f"{project.name}/{session.filename}"

            # 原始转写
            if session.raw_transcript:
                zip_file.writestr(
                    f"{session_dir}/transcript.txt",
                    session.raw_transcript
                )

            # 提取信息
            if session.extracted_info:
                zip_file.writestr(
                    f"{session_dir}/extracted.json",
                    session.extracted_info
                )

            # 补充信息
            if session.supplementary_info:
                zip_file.writestr(
                    f"{session_dir}/supplementary.txt",
                    session.supplementary_info
                )

            # 最终内容
            if session.final_content:
                zip_file.writestr(
                    f"{session_dir}/final.md",
                    session.final_content
                )

    zip_buffer.seek(0)

    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={
            "Content-Disposition": f"attachment; filename={project.name}_export.zip"
        }
    )


# ===== 合并创建项目和提纲 =====


def parse_markdown_with_title(content: str) -> tuple[str, OutlineContent]:
    """
    解析 Markdown 内容，提取标题作为项目名，其余作为提纲
    Returns: (project_name, outline_content)
    """
    lines = content.strip().split("\n")
    project_name = ""
    sections = []
    current_section = None

    for line in lines:
        line = line.strip()
        if not line:
            continue

        # 匹配标题
        header_match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if header_match:
            header_level = len(header_match.group(1))
            title = header_match.group(2)

            # 第一个一级标题作为项目名
            if header_level == 1 and not project_name:
                # 移除常见的后缀词
                project_name = title
                for suffix in ["提纲", "大纲", "问卷", "调研提纲", "访谈提纲"]:
                    if project_name.endswith(suffix):
                        project_name = project_name[:-len(suffix)].strip()
                        break
                continue

            # 其他标题作为板块
            if current_section:
                sections.append(current_section)

            current_section = OutlineSection(
                id=f"s{len(sections) + 1}",
                title=title,
                questions=[]
            )

        elif current_section:
            # 匹配列表项
            list_match = re.match(r"^[-*]\s+(.+)$|^\d+\.\s+(.+)$", line)
            if list_match:
                question_text = list_match.group(1) or list_match.group(2)
                current_section.questions.append(question_text)

    # 保存最后一个板块
    if current_section:
        sections.append(current_section)

    # 如果没有提取到项目名，使用默认值
    if not project_name:
        project_name = "新建项目"

    return project_name, OutlineContent(sections=sections)


@router.post("/import-with-outline", response_model=ProjectResponse)
async def import_project_with_outline(
    request: ProjectWithOutlineImport,
    db: AsyncSession = Depends(get_db),
):
    """
    同时创建项目和提纲
    - 从 Markdown 内容解析标题作为项目名
    - 其余内容作为提纲
    """
    # 解析 Markdown
    extracted_name, outline_content = parse_markdown_with_title(request.markdown_content)

    # 使用用户提供的项目名（如果有），否则使用解析出的名称
    project_name = request.project_name or extracted_name

    # 创建提纲
    outline_id = str(uuid4())
    outline = OutlineDB(
        id=outline_id,
        name=f"{project_name}提纲",
        content=outline_content.model_dump_json(),
    )
    db.add(outline)

    # 创建项目
    project_id = str(uuid4())
    project = ProjectDB(
        id=project_id,
        name=project_name,
        description=request.description or "",
        outline_id=outline_id,
    )
    db.add(project)

    await db.commit()
    await db.refresh(project)
    await db.refresh(outline)

    return ProjectResponse(
        id=project.id,
        name=project.name,
        description=project.description,
        outline_id=project.outline_id,
        table_structure_prompt=project.table_structure_prompt or "",
        summary_table=project.summary_table or "",
        outline=OutlineResponse(
            id=outline.id,
            name=outline.name,
            content=outline_content,
            created_at=outline.created_at,
            updated_at=outline.updated_at,
        ),
        session_count=0,
        created_at=project.created_at,
        updated_at=project.updated_at,
    )
